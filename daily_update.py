#!/usr/bin/env python3
"""
Daily Update Pipeline for Cek Barkot Tembakau
=============================================
Otomasi pembaruan data bal tembakau harian:
1. Membaca file Excel terbaru di folder 'Laporan Grade Induk'
2. Meregenerasi buku grade induk rekap (generate_laporan_grade_induk.py)
3. Memperbarui seed_data.json & seed_data.js
4. Memperbarui supabase_schema.sql
5. Bulk upsert langsung ke Supabase Cloud REST API
6. Memperbarui currentDate di app.js
7. Memperbarui subtitle tanggal & bump versi cache di index.html
8. Git add, commit, dan push ke origin main
"""

import os
import sys
import re
import json
import subprocess
import urllib.request
import urllib.error
from datetime import datetime
import openpyxl

# Set output encoding to UTF-8
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LAPORAN_DIR = os.path.join(BASE_DIR, "Laporan Grade Induk")

SUPABASE_URL = 'https://jrpklibocgicubevyshm.supabase.co'
SUPABASE_KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImpycGtsaWJvY2dpY3ViZXZ5c2htIiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODc4NTA3NjUsImV4cCI6MjEwMzQyNjc2NX0.xGoel8SNa2v9DcZBYwKcmjzGF7j6LJ-OQkr919JyYSc'

BULAN_MAP = {
    '1': 'JANUARI', '2': 'FEBRUARI', '3': 'MARET', '4': 'APRIL',
    '5': 'MEI', '6': 'JUNI', '7': 'JULI', '8': 'AGUSTUS',
    '9': 'SEPTEMBER', '10': 'OKTOBER', '11': 'NOVEMBER', '12': 'DESEMBER',
    'jan': 'JANUARI', 'feb': 'FEBRUARI', 'mar': 'MARET', 'apr': 'APRIL',
    'mei': 'MEI', 'may': 'MEI', 'jun': 'JUNI', 'jul': 'JULI',
    'agu': 'AGUSTUS', 'aug': 'AGUSTUS', 'sep': 'SEPTEMBER', 'sept': 'SEPTEMBER',
    'okt': 'OKTOBER', 'oct': 'OKTOBER', 'nov': 'NOVEMBER', 'des': 'DESEMBER', 'dec': 'DESEMBER'
}

def parse_excel_date(fpath):
    """Membaca tanggal, short date, dan hari dari file Excel grade harian."""
    wb = openpyxl.load_workbook(fpath, data_only=True)
    sheet = wb.active
    
    # Periksa baris 2: contoh 'TANGGAL: 23-9          HARI: RABU          TAHUN: 2026'
    r2_val = str(sheet.cell(row=2, column=1).value or '')
    
    tgl_match = re.search(r'TANGGAL:\s*(\d{1,2})-(\d{1,2})', r2_val, re.IGNORECASE)
    hari_match = re.search(r'HARI:\s*([A-Za-z]+)', r2_val, re.IGNORECASE)
    tahun_match = re.search(r'TAHUN:\s*(\d{4})', r2_val, re.IGNORECASE)
    
    tahun = tahun_match.group(1) if tahun_match else '2026'
    
    if tgl_match:
        d = int(tgl_match.group(1))
        m = int(tgl_match.group(2))
        iso_date = f"{tahun}-{m:02d}-{d:02d}"
        tgl_short = f"{d}-{m}"
        hari = hari_match.group(1).upper() if hari_match else "HARI"
        return iso_date, tgl_short, hari
    
    # Fallback dari nama file: contoh 'grade 23 sep.xlsx'
    fname = os.path.basename(fpath).lower()
    fn_match = re.search(r'grade\s+(?:tgl\s+)?(\d{1,2})\s*([a-z]+)', fname)
    if fn_match:
        d = int(fn_match.group(1))
        bln_str = fn_match.group(2)[:3]
        m = 9 if 'sep' in bln_str else (8 if 'agu' in bln_str else 9)
        iso_date = f"{tahun}-{m:02d}-{d:02d}"
        tgl_short = f"{d}-{m}"
        # Estimasi hari dari tanggal
        dt_obj = datetime.strptime(iso_date, "%Y-%m-%d")
        hari_names = ["SENIN", "SELASA", "RABU", "KAMIS", "JUMAT", "SABTU", "MINGGU"]
        hari = hari_names[dt_obj.weekday()]
        return iso_date, tgl_short, hari
        
    return None, None, None

def extract_items_from_excel(fpath, iso_date):
    """Mengekstrak baris bal dari file Excel."""
    wb = openpyxl.load_workbook(fpath, data_only=True)
    sheet = wb.active
    items = []
    
    for r in range(5, sheet.max_row + 1):
        no_gud = sheet.cell(row=r, column=1).value
        grade1 = sheet.cell(row=r, column=2).value
        grade2 = sheet.cell(row=r, column=3).value
        barkot = sheet.cell(row=r, column=4).value
        kg = sheet.cell(row=r, column=5).value
        ket = sheet.cell(row=r, column=6).value
        
        if any(v is not None for v in [no_gud, grade1, grade2, barkot, kg, ket]):
            b_str = str(barkot).strip() if barkot is not None else ''
            if b_str and b_str.isdigit():
                try:
                    ng_val = int(no_gud)
                except Exception:
                    ng_val = no_gud
                try:
                    kg_val = float(kg) if kg is not None else None
                except Exception:
                    kg_val = None
                items.append({
                    'tanggal': iso_date,
                    'no_gud': ng_val,
                    'grade': str(grade1).strip() if grade1 is not None else '',
                    'barkot': b_str,
                    'kg': kg_val,
                    'is_done': False
                })
    
    items.sort(key=lambda x: x['no_gud'] if isinstance(x['no_gud'], int) else 999999)
    return items

def update_generate_script(iso_date, tgl_short, hari, fname):
    """Memperbarui generate_laporan_grade_induk.py jika entri belum ada."""
    gen_path = os.path.join(BASE_DIR, "generate_laporan_grade_induk.py")
    with open(gen_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    entry_line = f'    ("{iso_date}", "{tgl_short}", "{hari}", "{fname}"),'
    if fname in content or iso_date in content:
        print(f"  [i] Entri {fname} sudah ada di generate_laporan_grade_induk.py")
        return
    
    # Sisipkan sebelum akhir list files
    target = ']\n\nall_items = []'
    replacement = f'{entry_line}\n]\n\nall_items = []'
    if target in content:
        content = content.replace(target, replacement)
    else:
        # Fallback regex
        content = re.sub(r'(\s*\(\"[^\"]+\",\s*\"[^\"]+\",\s*\"[^\"]+\",\s*\"[^\"]+\"\),\s*\n)(\])',
                         r'\1' + entry_line + r'\n\2', content)
    
    with open(gen_path, 'w', encoding='utf-8') as f:
        f.write(content)
    print(f"  [+] Ditambahkan entri ke generate_laporan_grade_induk.py: {entry_line.strip()}")

def run_generate_script():
    """Menjalankan generate_laporan_grade_induk.py untuk meregenerasi file Excel rekap."""
    print("  [*] Meregenerasi file Excel rekap buku grade induk...")
    gen_path = os.path.join(BASE_DIR, "generate_laporan_grade_induk.py")
    res = subprocess.run([sys.executable, gen_path], capture_output=True, text=True, cwd=BASE_DIR)
    if res.returncode != 0:
        print(f"  [!] Gagal menjalankan generate script: {res.stderr}")
    else:
        print("  [✓] File Excel rekap berhasil diperbarui.")

def update_seed_data(iso_date, items):
    """Memperbarui seed_data.json dan seed_data.js."""
    json_path = os.path.join(BASE_DIR, "seed_data.json")
    js_path = os.path.join(BASE_DIR, "seed_data.js")
    
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    data[iso_date] = items
    
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write('\n')
        
    with open(js_path, 'w', encoding='utf-8') as f:
        f.write('const SEED_DATA = ')
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write(';\n')
        
    print(f"  [✓] seed_data.json & seed_data.js diperbarui ({len(items)} bal).")

def update_supabase_schema(iso_date, items):
    """Memperbarui supabase_schema.sql dengan total baru dan klausa INSERT."""
    schema_path = os.path.join(BASE_DIR, "supabase_schema.sql")
    with open(schema_path, 'r', encoding='utf-8') as f:
        sql = f.read()
    
    # Hitung total bal dari seed_data.json
    json_path = os.path.join(BASE_DIR, "seed_data.json")
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    total_bal = sum(len(v) for v in data.values())
    
    # Update header comment
    sql = re.sub(
        r'-- 5\. Data Awal Bawaan \(Tanggal 2026-08-21 s/d \d{4}-\d{2}-\d{2} - Total \d+ Bal\)',
        f'-- 5. Data Awal Bawaan (Tanggal 2026-08-21 s/d {iso_date} - Total {total_bal} Bal)',
        sql
    )
    
    # Cek apakah tanggal sudah ada di INSERT
    if f"('{iso_date}'," in sql:
        print(f"  [i] Record {iso_date} sudah ada di supabase_schema.sql")
        with open(schema_path, 'w', encoding='utf-8') as f:
            f.write(sql)
        return
    
    # Ganti titik koma terakhir menjadi koma dan tambahkan baris baru
    # Cari baris insert terakhir: ('YYYY-MM-DD', ..., ...);
    last_semi_match = re.search(r"(\('[^']+',\s*\d+,\s*[^;]+);\s*$", sql.rstrip())
    if last_semi_match:
        last_item_str = last_semi_match.group(1)
        new_lines = [f"{last_item_str},"]
        for idx, it in enumerate(items):
            term = ';' if idx == len(items) - 1 else ','
            grade_val = f"'{it['grade']}'" if it['grade'] else 'NULL'
            barkot_val = f"'{it['barkot']}'" if it['barkot'] else 'NULL'
            kg_val = f"{it['kg']:.1f}" if it['kg'] is not None else 'NULL'
            is_done_val = 'true' if it.get('is_done') else 'false'
            line = f"('{it['tanggal']}', {it['no_gud']}, {grade_val}, {barkot_val}, {kg_val}, {is_done_val}){term}"
            new_lines.append(line)
        
        sql = sql[:last_semi_match.start()] + '\n'.join(new_lines) + '\n'
        with open(schema_path, 'w', encoding='utf-8') as f:
            f.write(sql)
        print(f"  [✓] supabase_schema.sql diperbarui (Total {total_bal} bal).")

def sync_to_supabase(items):
    """Bulk upsert langsung ke REST API Supabase Cloud."""
    if not items:
        return
    
    iso_date = items[0]['tanggal']
    print(f"  [*] Mengirim {len(items)} record {iso_date} ke Supabase Cloud REST API...")
    payload = []
    for it in items:
        payload.append({
            'tanggal': it['tanggal'],
            'no_gud': it['no_gud'],
            'grade': it['grade'] or None,
            'barkot': it['barkot'] or None,
            'kg': it['kg'],
            'is_done': False
        })
        
    url = f"{SUPABASE_URL}/rest/v1/barkot_data?on_conflict=tanggal,no_gud"
    headers = {
        'apikey': SUPABASE_KEY,
        'Authorization': f'Bearer {SUPABASE_KEY}',
        'Content-Type': 'application/json',
        'Prefer': 'resolution=merge-duplicates,return=representation'
    }
    
    req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers=headers, method='POST')
    try:
        with urllib.request.urlopen(req) as resp:
            print(f"  [✓] Supabase Cloud merespons HTTP {resp.status}. Data tersimpan.")
    except Exception as e:
        print(f"  [!] Peringatan sync Supabase: {e}")
        
    # Verifikasi jumlah record
    try:
        v_url = f"{SUPABASE_URL}/rest/v1/barkot_data?select=count&tanggal=eq.{iso_date}"
        v_req = urllib.request.Request(v_url, headers={
            'apikey': SUPABASE_KEY,
            'Authorization': f'Bearer {SUPABASE_KEY}',
            'Prefer': 'count=exact'
        })
        with urllib.request.urlopen(v_req) as v_resp:
            body = json.loads(v_resp.read().decode('utf-8'))
            print(f"  [✓] Verifikasi Cloud: Total {body[0]['count']} bal aktif di tabel Supabase.")
    except Exception as e:
        print(f"  [!] Gagal memverifikasi count Supabase: {e}")

def update_frontend(iso_date):
    """Memperbarui currentDate di app.js dan bump versi cache di index.html."""
    app_path = os.path.join(BASE_DIR, "app.js")
    idx_path = os.path.join(BASE_DIR, "index.html")
    
    # 1. Update app.js
    with open(app_path, 'r', encoding='utf-8') as f:
        app_js = f.read()
    app_js = re.sub(r"let currentDate = '[^']+';", f"let currentDate = '{iso_date}';", app_js)
    with open(app_path, 'w', encoding='utf-8') as f:
        f.write(app_js)
        
    # 2. Update index.html (subtitle tanggal & bump cache version)
    with open(idx_path, 'r', encoding='utf-8') as f:
        html = f.read()
        
    # Format tanggal: misal 2026-09-23 -> "23 SEPTEMBER"
    parts = iso_date.split('-')
    d_num = int(parts[2])
    m_num = str(int(parts[1]))
    bln_name = BULAN_MAP.get(m_num, 'SEPTEMBER')
    date_label = f"{d_num} {bln_name}"
    
    html = re.sub(
        r'<span class="display-subtitle" id="dateDisplayLabel">[^<]+</span>',
        f'<span class="display-subtitle" id="dateDisplayLabel">{date_label}</span>',
        html
    )
    
    # Cari versi saat ini (misal v=6.5) dan naikkan 0.1
    v_match = re.search(r'href="style\.css\?v=([0-9\.]+)"', html)
    if v_match:
        cur_v = float(v_match.group(1))
        new_v = round(cur_v + 0.1, 1)
        new_v_str = f"v={new_v}"
        old_v_str = f"v={v_match.group(1)}"
        
        html = html.replace(f"style.css?{old_v_str}", f"style.css?{new_v_str}")
        html = html.replace(f"seed_data.js?{old_v_str}", f"seed_data.js?{new_v_str}")
        html = html.replace(f"app.js?{old_v_str}", f"app.js?{new_v_str}")
        print(f"  [✓] Cache version dinaikkan dari {old_v_str} -> {new_v_str}")
    
    with open(idx_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"  [✓] index.html & app.js diperbarui (Default: {date_label}).")
    return new_v_str if v_match else "v=latest"

def git_commit_and_push(iso_date, items, cache_ver):
    """Melakukan git add, commit, dan push."""
    print("  [*] Menjalankan Git Commit & Push...")
    min_ng = min(x['no_gud'] for x in items) if items else 0
    max_ng = max(x['no_gud'] for x in items) if items else 0
    
    # Format tanggal deskriptif
    parts = iso_date.split('-')
    d_num = int(parts[2])
    m_num = str(int(parts[1]))
    bln_name = BULAN_MAP.get(m_num, 'September').capitalize()
    thn = parts[0]
    
    commit_msg = (
        f"feat: tambah data bal tembakau tanggal {d_num} {bln_name} {thn} "
        f"({len(items)} bal, No Gud {min_ng}-{max_ng}), sync Supabase, update rekap excel, dan bump cache version {cache_ver}"
    )
    
    subprocess.run(["git", "add", "."], cwd=BASE_DIR, check=True)
    
    # Periksa apakah ada perubahan yang di-stage
    status_res = subprocess.run(["git", "status", "--porcelain"], cwd=BASE_DIR, capture_output=True, text=True)
    if not status_res.stdout.strip():
        print("  [i] Tidak ada perubahan berkas untuk di-commit.")
        return
        
    subprocess.run(["git", "commit", "-m", commit_msg], cwd=BASE_DIR, check=True)
    print(f"  [✓] Commit berhasil: '{commit_msg}'")
    
    push_res = subprocess.run(["git", "push", "origin", "main"], cwd=BASE_DIR, capture_output=True, text=True)
    if push_res.returncode == 0:
        print("  [✓] Git push berhasil ke origin main!")
        print("  👉 Web live di: https://grndnstnmlk.github.io/cekbarkot/")
    else:
        print(f"  [!] Git push error: {push_res.stderr}")

def find_target_excel():
    """Mencari file excel terbaru di folder Laporan Grade Induk yang belum ada di seed_data.json."""
    json_path = os.path.join(BASE_DIR, "seed_data.json")
    existing_dates = set()
    if os.path.exists(json_path):
        with open(json_path, 'r', encoding='utf-8') as f:
            existing_dates = set(json.load(f).keys())
            
    excel_files = [f for f in os.listdir(LAPORAN_DIR) if f.endswith('.xlsx') and not f.startswith('~$') and not f.startswith('buku_') and not f.startswith('laporan_')]
    
    # Urutkan berdasarkan waktu modifikasi terbaru
    excel_files.sort(key=lambda x: os.path.getmtime(os.path.join(LAPORAN_DIR, x)), reverse=True)
    
    for f in excel_files:
        fpath = os.path.join(LAPORAN_DIR, f)
        iso_date, tgl_short, hari = parse_excel_date(fpath)
        if iso_date and iso_date not in existing_dates:
            return fpath, iso_date, tgl_short, hari
            
    # Jika semua sudah ada, ambil yang paling baru
    if excel_files:
        latest = excel_files[0]
        fpath = os.path.join(LAPORAN_DIR, latest)
        iso_date, tgl_short, hari = parse_excel_date(fpath)
        return fpath, iso_date, tgl_short, hari
        
    return None, None, None, None

def run_pipeline(target_file=None):
    print("=" * 60)
    print("🌿 MEMULAI DAILY UPDATE PIPELINE CEK BARKOT")
    print("=" * 60)
    
    if target_file:
        fpath = target_file if os.path.isabs(target_file) else os.path.join(LAPORAN_DIR, target_file)
        iso_date, tgl_short, hari = parse_excel_date(fpath)
    else:
        fpath, iso_date, tgl_short, hari = find_target_excel()
        
    if not fpath or not os.path.exists(fpath):
        print("[-] Tidak ditemukan file Excel yang cocok di folder Laporan Grade Induk.")
        sys.exit(1)
        
    fname = os.path.basename(fpath)
    print(f"[*] Target File : {fname}")
    print(f"[*] Tanggal     : {iso_date} ({hari}, {tgl_short})")
    
    # 1. Ekstrak data bal
    items = extract_items_from_excel(fpath, iso_date)
    total_kg = sum(x['kg'] for x in items if x['kg'])
    print(f"[*] Jumlah Bal  : {len(items)} bal (No Gud {min(x['no_gud'] for x in items)} s/d {max(x['no_gud'] for x in items)})")
    print(f"[*] Total Berat : {total_kg:.2f} Kg")
    
    # 2. Update generate_laporan_grade_induk.py & run
    update_generate_script(iso_date, tgl_short, hari, fname)
    run_generate_script()
    
    # 3. Update seed_data.json & seed_data.js
    update_seed_data(iso_date, items)
    
    # 4. Update supabase_schema.sql
    update_supabase_schema(iso_date, items)
    
    # 5. Sync to Supabase Cloud
    sync_to_supabase(items)
    
    # 6. Update app.js & index.html
    cache_ver = update_frontend(iso_date)
    
    # 7. Git commit & push
    git_commit_and_push(iso_date, items, cache_ver)
    
    print("=" * 60)
    print("🎉 PEMBARUAN SELESAI DENGAN SUKSES!")
    print("=" * 60)

if __name__ == '__main__':
    specified_file = sys.argv[1] if len(sys.argv) > 1 else None
    run_pipeline(specified_file)
