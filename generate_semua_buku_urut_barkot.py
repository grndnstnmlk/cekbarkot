#!/usr/bin/env python3
"""
Generator Buku Grade Induk Urut Barkot (Siap Print)
===================================================
Menghasilkan berkas Excel Buku Grade Induk per tanggal yang:
1. Diurutkan ascending berdasarkan nomor seri BARKOT (barcode stiker fisik).
2. Diformat rapi dan presisi untuk pencetakan dokumen fisik kertas A4:
   - Sheet 1: "Cetak 1 Halaman (A4 Pas)" -> Fit to 1 page A4 portrait.
   - Sheet 2: "Cetak Standar (Font Besar)" -> Font size 14, repeating headers $3:$4, auto page height.
3. Disimpan terpusat dan rapi di dalam folder:
   `Buku Grade Induk Urut Barkot/`
4. Menghasilkan berkas master index rekap cetak:
   `_REKAP_DAFTAR_PRINT_SEMUA_TANGGAL.xlsx`
"""

import os
import sys
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LAPORAN_DIR = os.path.join(BASE_DIR, "Laporan Grade Induk")
OUTPUT_DIR = os.path.join(BASE_DIR, "Buku Grade Induk Urut Barkot")

MONTH_NAMES_ID = {
    '01': 'Jan', '02': 'Feb', '03': 'Mar', '04': 'Apr',
    '05': 'Mei', '06': 'Jun', '07': 'Jul', '08': 'Ags',
    '09': 'Sep', '10': 'Okt', '11': 'Nov', '12': 'Des'
}

MONTH_FULL_ID = {
    '01': 'Januari', '02': 'Februari', '03': 'Maret', '04': 'April',
    '05': 'Mei', '06': 'Juni', '07': 'Juli', '08': 'Agustus',
    '09': 'September', '10': 'Oktober', '11': 'November', '12': 'Desember'
}

# Daftar seluruh berkas grade harian (Tanggal, tgl_short, hari, nama file sumber)
FILES_MASTER = [
    ("2026-08-21", "21-8", "JUMAT", "grade tgl 21.xlsx"),
    ("2026-08-22", "22-8", "SABTU", "grade tgl 22.xlsx"),
    ("2026-08-23", "23-8", "MINGGU", "grade tgl 23.xlsx"),
    ("2026-08-24", "24-8", "SENIN", "grade tgl 24.xlsx"),
    ("2026-08-26", "26-8", "RABU", "grade induk tgl 26.xlsx"),
    ("2026-08-27", "27-8", "KAMIS", "buku_grade_induk_27-8-2026.xlsx"),
    ("2026-08-28", "28-8", "JUMAT", "grade tgl 28.xlsx"),
    ("2026-08-29", "29-8", "SABTU", "grade lengkap tgl 29.xlsx"),
    ("2026-08-30", "30-8", "MINGGU", "grade tgl 30.xlsx"),
    ("2026-08-31", "31-8", "SENIN", "grade tgl 31.xlsx"),
    ("2026-09-01", "1-9", "SELASA", "grade_tgl_1_september_sorted.xlsx"),
    ("2026-09-02", "2-9", "RABU", "Grade 2 sep.xlsx"),
    ("2026-09-03", "3-9", "KAMIS", "grade 3 sep.xlsx"),
    ("2026-09-04", "4-9", "JUMAT", "grade 4 sep.xlsx"),
    ("2026-09-05", "5-9", "SABTU", "grade 5 sep.xlsx"),
    ("2026-09-06", "6-9", "MINGGU", "grade 6 september.xlsx"),
    ("2026-09-07", "7-9", "SENIN", "grade 7 september.xlsx"),
    ("2026-09-08", "8-9", "SELASA", "grade 8 sep.xlsx"),
    ("2026-09-09", "9-9", "RABU", "grade 9 sep.xlsx"),
    ("2026-09-10", "10-9", "KAMIS", "grade 10 sep.xlsx"),
    ("2026-09-11", "11-9", "JUMAT", "grade 11 sep.xlsx"),
    ("2026-09-12", "12-9", "SABTU", "grade 12 sep.xlsx"),
    ("2026-09-13", "13-9", "MINGGU", "grade 13 sep.xlsx"),
    ("2026-09-14", "14-9", "SENIN", "grade 14 sep.xlsx"),
    ("2026-09-17", "17-9", "KAMIS", "grade 17 september.xlsx"),
    ("2026-09-18", "18-9", "JUMAT", "grade 18 sep.xlsx"),
    ("2026-09-19", "19-9", "SABTU", "grade 19 september.xlsx"),
    ("2026-09-20", "20-9", "MINGGU", "grade 20 sep.xlsx"),
    ("2026-09-21", "21-9", "SENIN", "grade 21 sep.xlsx"),
    ("2026-09-22", "22-9", "SELASA", "grade 22 sep.xlsx"),
    ("2026-09-23", "23-9", "RABU", "grade 23 sep.xlsx"),
    ("2026-09-24", "24-9", "KAMIS", "grade 24 sep.xlsx"),
    ("2026-09-25", "25-9", "JUMAT", "grade 25 sep.xlsx"),
    ("2026-09-26", "26-9", "SABTU", "grade 26 sep.xlsx"),
    ("2026-09-27", "27-9", "MINGGU", "grade 27 sep.xlsx"),
    ("2026-09-28", "28-9", "SENIN", "grade 28 sep.xlsx"),
    ("2026-09-29", "29-9", "SELASA", "grade 29 sep.xlsx"),
    ("2026-09-30", "30-9", "RABU", "grade 30 sep.xlsx"),
    ("2026-10-01", "1-10", "KAMIS", "grade 1 okt.xlsx"),
    ("2026-10-02", "2-10", "JUMAT", "grade 2 okt.xlsx"),
    ("2026-10-03", "3-10", "SABTU", "grade 3 okt.xlsx"),
    ("2026-10-04", "4-10", "MINGGU", "grade 4 okt.xlsx"),
    ("2026-10-05", "5-10", "SENIN", "grade 5 okt.xlsx"),
    ("2026-10-06", "6-10", "SELASA", "grade 6 okt.xlsx"),
]

def make_filename(dt_str):
    """Menghasilkan nama file yang terurut rapi secara kronologis di Windows Explorer."""
    parts = dt_str.split('-')
    y, m, d = parts[0], parts[1], parts[2]
    m_abbr = MONTH_NAMES_ID.get(m, m).lower()
    return f"buku_grade_induk_{y}-{m}-{d}_tgl_{d}_{m_abbr}_urut_barkot.xlsx"

def extract_and_sort_items(fpath):
    """Mengekstrak baris bal dari file sumber dan mengurutkan secara numerik berdasarkan nomor seri barkot."""
    wb_src = openpyxl.load_workbook(fpath, data_only=True)
    sheet_src = wb_src.active
    
    items = []
    for r in range(5, sheet_src.max_row + 1):
        no_gud = sheet_src.cell(row=r, column=1).value
        grade1 = sheet_src.cell(row=r, column=2).value
        grade2 = sheet_src.cell(row=r, column=3).value
        barkot = sheet_src.cell(row=r, column=4).value
        kg = sheet_src.cell(row=r, column=5).value
        ket = sheet_src.cell(row=r, column=6).value
        
        if any(v is not None for v in [no_gud, grade1, grade2, barkot, kg, ket]):
            b_str = str(barkot).strip() if barkot is not None else ""
            if b_str and b_str.isdigit():
                try:
                    ng_int = int(no_gud)
                except Exception:
                    ng_int = no_gud
                
                kg_val = None
                if kg is not None:
                    try:
                        kg_val = float(kg)
                    except Exception:
                        kg_val = kg
                        
                items.append({
                    "no_gud": ng_int,
                    "grade1": str(grade1).strip() if grade1 is not None else "",
                    "grade2": str(grade2).strip() if grade2 is not None else "",
                    "barkot": b_str,
                    "kg": kg_val,
                    "ket": str(ket).strip() if ket is not None else ""
                })
                
    # Urutkan ascending berdasarkan nomor barkot (numerik: 30140 < 30141 < ... < 578206)
    items.sort(key=lambda x: int(x["barkot"]))
    return items

def build_daily_workbook(dt_str, tgl_short, hari, items):
    """Membangun workbook Excel dengan 2 sheet cetak siap pakai (A4 Pas & Font Besar)."""
    wb = openpyxl.Workbook()
    
    parts = dt_str.split('-')
    tahun = parts[0]
    
    # Styles
    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    bottom_thin_border = Border(bottom=thin_side)
    align_center = Alignment(horizontal="center", vertical="center")
    
    start_row = 5
    last_data_row = start_row + len(items) - 1 if items else start_row
    total_row = last_data_row + 1
    
    # ========================================================
    # SHEET 1: CETAK 1 HALAMAN (A4 PAS - MARGIN MAKSIMAL)
    # ========================================================
    ws1 = wb.active
    ws1.title = "Cetak 1 Halaman (A4 Pas)"
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
    ws1.page_margins.left = 0.25
    ws1.page_margins.right = 0.25
    ws1.page_margins.top = 0.30
    ws1.page_margins.bottom = 0.30
    ws1.page_margins.header = 0.15
    ws1.page_margins.footer = 0.15
    ws1.print_options.horizontalCentered = True
    ws1.print_options.verticalCentered = False
    ws1.sheet_properties.pageSetUpPr.fitToPage = True
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 1
    ws1.print_area = f"A1:F{total_row}"
    ws1.views.sheetView[0].showGridLines = True
    
    # Column Dimensions
    ws1.column_dimensions['A'].width = 11.0   # NO GUD
    ws1.column_dimensions['B'].width = 8.0    # GRADE 1
    ws1.column_dimensions['C'].width = 8.0    # GRADE 2
    ws1.column_dimensions['D'].width = 16.0   # BARKOT
    ws1.column_dimensions['E'].width = 8.5    # KG
    ws1.column_dimensions['F'].width = 16.0   # KET
    
    # Row Heights
    ws1.row_dimensions[1].height = 24.0
    ws1.row_dimensions[2].height = 15.0
    ws1.row_dimensions[3].height = 16.0
    ws1.row_dimensions[4].height = 16.0
    
    # Row 1: Judul
    ws1.merge_cells("A1:F1")
    c_a1 = ws1["A1"]
    c_a1.value = "BUKU GRADE INDUK"
    c_a1.font = Font(name="Arial", size=15, bold=True, color="000000")
    c_a1.alignment = align_center
    
    # Row 2: Subtitle
    ws1.merge_cells("A2:F2")
    c_a2 = ws1["A2"]
    c_a2.value = f"TANGGAL: {tgl_short}          HARI: {hari}          TAHUN: {tahun}"
    c_a2.font = Font(name="Arial", size=9, bold=True, color="000000")
    c_a2.alignment = align_center
    for col in range(1, 7):
        ws1.cell(row=2, column=col).border = bottom_thin_border
        
    # Row 3 & 4: Table Headers
    for rng, val in [("A3:A4", "NO GUD"), ("B3:C4", "GRADE"), ("D3:D4", "BARKOT"), ("E3:E4", "KG"), ("F3:F4", "KET")]:
        ws1.merge_cells(rng)
        top_left = ws1[rng.split(":")[0]]
        top_left.value = val
        
    for r in [3, 4]:
        for c in range(1, 7):
            cell = ws1.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = fill_navy
            cell.alignment = align_center
            cell.border = thin_border
            
    # Data Rows
    font_data_s1 = Font(name="Arial", size=11, bold=False, color="000000")
    for idx, item in enumerate(items):
        curr_row = start_row + idx
        ws1.row_dimensions[curr_row].height = 14.5
        
        c_a = ws1.cell(row=curr_row, column=1, value=item["no_gud"])
        c_b = ws1.cell(row=curr_row, column=2, value=item["grade1"])
        c_c = ws1.cell(row=curr_row, column=3, value=item["grade2"] if item["grade2"] else None)
        c_d = ws1.cell(row=curr_row, column=4, value=str(item["barkot"]))
        c_d.number_format = '@'
        c_e = ws1.cell(row=curr_row, column=5, value=item["kg"])
        if item["kg"] is not None and isinstance(item["kg"], (int, float)):
            c_e.number_format = '0.0'
        c_f = ws1.cell(row=curr_row, column=6, value=item["ket"] if item["ket"] else None)
        
        for cell in [c_a, c_b, c_c, c_d, c_e, c_f]:
            cell.font = font_data_s1
            cell.alignment = align_center
            cell.border = thin_border
            
    # Total Row
    ws1.row_dimensions[total_row].height = 18.0
    ws1.merge_cells(f"A{total_row}:D{total_row}")
    cell_tot_lbl = ws1[f"A{total_row}"]
    cell_tot_lbl.value = f"TOTAL ({len(items)} BAL)"
    
    cell_tot_kg = ws1[f"E{total_row}"]
    cell_tot_kg.value = f"=SUM(E{start_row}:E{last_data_row})"
    cell_tot_kg.number_format = '0.0'
    
    cell_tot_ket = ws1[f"F{total_row}"]
    cell_tot_ket.value = "KG"
    
    for c in range(1, 7):
        cell = ws1.cell(row=total_row, column=c)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = thin_border

    # ========================================================
    # SHEET 2: CETAK STANDAR (FONT BESAR - MULTI HALAMAN)
    # ========================================================
    ws2 = wb.create_sheet(title="Cetak Standar (Font Besar)")
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.page_setup.orientation = ws2.ORIENTATION_PORTRAIT
    ws2.page_margins.left = 0.25
    ws2.page_margins.right = 0.25
    ws2.page_margins.top = 0.35
    ws2.page_margins.bottom = 0.35
    ws2.page_margins.header = 0.15
    ws2.page_margins.footer = 0.15
    ws2.print_options.horizontalCentered = True
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.print_title_rows = "$3:$4"
    ws2.print_area = f"A1:F{total_row}"
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.column_dimensions['A'].width = 13.0
    ws2.column_dimensions['B'].width = 9.0
    ws2.column_dimensions['C'].width = 10.0
    ws2.column_dimensions['D'].width = 20.0
    ws2.column_dimensions['E'].width = 9.0
    ws2.column_dimensions['F'].width = 20.0
    
    ws2.row_dimensions[1].height = 31.5
    ws2.row_dimensions[2].height = 21.75
    ws2.row_dimensions[3].height = 21.0
    ws2.row_dimensions[4].height = 21.0
    
    ws2.merge_cells("A1:F1")
    ws2["A1"].value = "BUKU GRADE INDUK"
    ws2["A1"].font = Font(name="Arial", size=20, bold=True, color="000000")
    ws2["A1"].alignment = align_center
    
    ws2.merge_cells("A2:F2")
    ws2["A2"].value = f"TANGGAL: {tgl_short}          HARI: {hari}          TAHUN: {tahun}"
    ws2["A2"].font = Font(name="Arial", size=10, bold=True, color="000000")
    ws2["A2"].alignment = align_center
    for col in range(1, 7):
        ws2.cell(row=2, column=col).border = bottom_thin_border
        
    for rng, val in [("A3:A4", "NO GUD"), ("B3:C4", "GRADE"), ("D3:D4", "BARKOT"), ("E3:E4", "KG"), ("F3:F4", "KET")]:
        ws2.merge_cells(rng)
        top_left = ws2[rng.split(":")[0]]
        top_left.value = val
        
    for r in [3, 4]:
        for c in range(1, 7):
            cell = ws2.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
            cell.fill = fill_navy
            cell.alignment = align_center
            cell.border = thin_border
            
    font_data_s2 = Font(name="Arial", size=14, bold=False, color="000000")
    for idx, item in enumerate(items):
        curr_row = start_row + idx
        ws2.row_dimensions[curr_row].height = 20.25
        
        c_a = ws2.cell(row=curr_row, column=1, value=item["no_gud"])
        c_b = ws2.cell(row=curr_row, column=2, value=item["grade1"])
        c_c = ws2.cell(row=curr_row, column=3, value=item["grade2"] if item["grade2"] else None)
        c_d = ws2.cell(row=curr_row, column=4, value=str(item["barkot"]))
        c_d.number_format = '@'
        c_e = ws2.cell(row=curr_row, column=5, value=item["kg"])
        if item["kg"] is not None and isinstance(item["kg"], (int, float)):
            c_e.number_format = '0.0'
        c_f = ws2.cell(row=curr_row, column=6, value=item["ket"] if item["ket"] else None)
        
        for cell in [c_a, c_b, c_c, c_d, c_e, c_f]:
            cell.font = font_data_s2
            cell.alignment = align_center
            cell.border = thin_border
            
    ws2.row_dimensions[total_row].height = 22.0
    ws2.merge_cells(f"A{total_row}:D{total_row}")
    cell_tot_lbl2 = ws2[f"A{total_row}"]
    cell_tot_lbl2.value = f"TOTAL ({len(items)} BAL)"
    
    cell_tot_kg2 = ws2[f"E{total_row}"]
    cell_tot_kg2.value = f"=SUM(E{start_row}:E{last_data_row})"
    cell_tot_kg2.number_format = '0.0'
    
    cell_tot_ket2 = ws2[f"F{total_row}"]
    cell_tot_ket2.value = "KG"
    
    for c in range(1, 7):
        cell = ws2.cell(row=total_row, column=c)
        cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = thin_border
        
    return wb

def build_summary_index_workbook(summary_rows):
    """Membangun berkas master index checklist rekap pencetakan seluruh tanggal."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daftar Rekap Print Barkot"
    
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "$3:$4"
    ws.views.sheetView[0].showGridLines = True
    
    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_gold = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # Title
    ws.merge_cells("A1:H1")
    ws["A1"].value = "DAFTAR BERKAS CETAK BUKU GRADE INDUK (URUT BARKOT)"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E78")
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 25.0
    
    ws.merge_cells("A2:H2")
    ws["A2"].value = "Folder: Buku Grade Induk Urut Barkot  |  Format Cetak: Kertas A4 Portrait (Siap Print)"
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")
    ws["A2"].alignment = align_center
    ws.row_dimensions[2].height = 16.0
    
    headers = [
        ("NO", 5.0),
        ("TANGGAL", 13.0),
        ("HARI", 10.0),
        ("JUMLAH BAL", 12.0),
        ("TOTAL KG", 11.0),
        ("RENTANG BARKOT", 22.0),
        ("RENTANG NO GUD", 18.0),
        ("NAMA BERKAS EXCEL", 45.0)
    ]
    
    ws.row_dimensions[3].height = 22.0
    for c_idx, (h_name, w) in enumerate(headers, 1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = w
        cell = ws.cell(row=3, column=c_idx, value=h_name)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = thin_border
        
    start_r = 4
    for idx, row in enumerate(summary_rows, 1):
        curr_r = start_r + idx - 1
        ws.row_dimensions[curr_r].height = 18.0
        
        ws.cell(row=curr_r, column=1, value=idx).alignment = align_center
        ws.cell(row=curr_r, column=2, value=row["tanggal"]).alignment = align_center
        ws.cell(row=curr_r, column=3, value=row["hari"]).alignment = align_center
        
        c_bal = ws.cell(row=curr_r, column=4, value=row["bal_count"])
        c_bal.alignment = align_center
        c_bal.number_format = '#,##0'
        
        c_kg = ws.cell(row=curr_r, column=5, value=row["total_kg"])
        c_kg.alignment = align_center
        if row["total_kg"] is not None:
            c_kg.number_format = '#,##0.0'
            
        ws.cell(row=curr_r, column=6, value=row["barkot_range"]).alignment = align_center
        ws.cell(row=curr_r, column=7, value=row["nogud_range"]).alignment = align_center
        ws.cell(row=curr_r, column=8, value=row["filename"]).alignment = align_left
        
        is_even = (idx % 2 == 0)
        for c in range(1, 9):
            cell = ws.cell(row=curr_r, column=c)
            cell.font = Font(name="Arial", size=10, bold=False)
            cell.border = thin_border
            if is_even:
                cell.fill = fill_zebra
                
    # Total row
    tot_r = start_r + len(summary_rows)
    ws.row_dimensions[tot_r].height = 22.0
    ws.merge_cells(f"A{tot_r}:C{tot_r}")
    ws[f"A{tot_r}"].value = "TOTAL KESELURUHAN"
    ws[f"A{tot_r}"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    ws[f"A{tot_r}"].alignment = align_center
    
    c_tot_bal = ws.cell(row=tot_r, column=4, value=f"=SUM(D{start_r}:D{tot_r-1})")
    c_tot_bal.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    c_tot_bal.alignment = align_center
    c_tot_bal.number_format = '#,##0'
    
    c_tot_kg = ws.cell(row=tot_r, column=5, value=f"=SUM(E{start_r}:E{tot_r-1})")
    c_tot_kg.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    c_tot_kg.alignment = align_center
    c_tot_kg.number_format = '#,##0.0'
    
    for c in range(1, 9):
        cell = ws.cell(row=tot_r, column=c)
        cell.fill = fill_navy
        cell.border = thin_border
        
    return wb

def generate_all_sorted_daily_files(target_date=None):
    """Menghasilkan berkas buku grade induk urut barkot untuk semua tanggal atau tanggal tertentu."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    targets = FILES_MASTER
    if target_date:
        targets = [x for x in FILES_MASTER if x[0] == target_date]
        if not targets:
            print(f"[!] Tanggal {target_date} tidak ditemukan di daftar FILES_MASTER.")
            return
            
    print("=" * 65)
    print("🖨️  MEMULAI GENERATOR BUKU GRADE INDUK URUT BARKOT (SIAP PRINT)")
    print(f"📁 Folder Tujuan : {OUTPUT_DIR}")
    print(f"📄 Jumlah Target  : {len(targets)} tanggal operasional")
    print("=" * 65)
    
    summary_rows = []
    
    for idx, (dt_str, tgl_short, hari, fname) in enumerate(targets, 1):
        src_fpath = os.path.join(LAPORAN_DIR, fname)
        if not os.path.exists(src_fpath):
            print(f"  [!] Berkas sumber tidak ditemukan: {fname}")
            continue
            
        items = extract_and_sort_items(src_fpath)
        if not items:
            print(f"  [!] Tidak ada data bal valid di {fname}")
            continue
            
        out_fname = make_filename(dt_str)
        out_fpath = os.path.join(OUTPUT_DIR, out_fname)
        
        # Bangun workbook 2 sheet
        wb = build_daily_workbook(dt_str, tgl_short, hari, items)
        wb.save(out_fpath)
        
        # Metadata ringkasan
        total_kg_val = sum(x["kg"] for x in items if x["kg"] is not None) if any(x["kg"] for x in items) else None
        min_ng = min(x["no_gud"] for x in items if isinstance(x["no_gud"], int))
        max_ng = max(x["no_gud"] for x in items if isinstance(x["no_gud"], int))
        barkot_min = items[0]["barkot"]
        barkot_max = items[-1]["barkot"]
        
        summary_rows.append({
            "tanggal": dt_str,
            "hari": hari,
            "bal_count": len(items),
            "total_kg": total_kg_val,
            "barkot_range": f"{barkot_min} s/d {barkot_max}",
            "nogud_range": f"No. Gud {min_ng} - {max_ng}",
            "filename": out_fname
        })
        
        kg_display = f"{total_kg_val:.1f} Kg" if total_kg_val else "- Kg"
        print(f"[{idx:02d}/{len(targets):02d}] ✓ {dt_str} ({hari}, {tgl_short}) -> {len(items)} bal | {kg_display} | Barkot: {barkot_min}..{barkot_max} -> '{out_fname}'")

    # Generate master index table
    if summary_rows:
        idx_wb = build_summary_index_workbook(summary_rows)
        idx_fpath = os.path.join(OUTPUT_DIR, "_REKAP_DAFTAR_PRINT_SEMUA_TANGGAL.xlsx")
        idx_wb.save(idx_fpath)
        print("=" * 65)
        print(f"✓ Master Index tersimpan: {idx_fpath}")
        print(f"🎉 SUKSES! Seluruh {len(summary_rows)} berkas siap cetak telah rapi di folder:")
        print(f"👉 {OUTPUT_DIR}")
        print("=" * 65)

if __name__ == '__main__':
    t_date = sys.argv[1] if len(sys.argv) > 1 else None
    generate_all_sorted_daily_files(t_date)
